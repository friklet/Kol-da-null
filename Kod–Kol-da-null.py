"""
Генератор презентации:
«XXI век: модернизация и новые стандарты в нефтепереработке»
10 слайдов, минималистичный дизайн, Times New Roman, подписи к рисункам.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ----------------------------------------------------------------------
# ЦВЕТОВАЯ ПАЛИТРА (нефтяная тематика)
# ----------------------------------------------------------------------
COLOR_DARK   = RGBColor(0x0B, 0x1F, 0x3A)   # тёмно-синий (фон титула)
COLOR_ACCENT = RGBColor(0xC8, 0x8A, 0x2B)   # золотистый (акцент)
COLOR_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_GRAY   = RGBColor(0x59, 0x59, 0x59)
COLOR_LIGHT  = RGBColor(0xF2, 0xF4, 0xF7)   # светлый фон контента

FONT = "Times New Roman"

# ----------------------------------------------------------------------
# ХЕЛПЕРЫ
# ----------------------------------------------------------------------
def add_title_bar(slide, text, prs):
    """Верхняя плашка с заголовком слайда."""
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.0)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_DARK
    bar.line.fill.background()

    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(1.0), prs.slide_width, Inches(0.06)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = COLOR_ACCENT
    accent.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.4), Inches(0.15), prs.slide_width - Inches(0.8), Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = COLOR_WHITE


def add_bullets(slide, items, left, top, width, height, size=18):
    """Маркированный список."""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)
        r = p.add_run()
        r.text = "• " + item
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.color.rgb = COLOR_DARK
    return tb


def add_image_placeholder(slide, left, top, width, height, caption=None, label="Рис."):
    """Рамка-заглушка под изображение + подпись снизу."""
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = COLOR_LIGHT
    box.line.color.rgb = COLOR_GRAY
    box.line.width = Pt(1)

    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "[ место для изображения ]"
    r.font.name = FONT
    r.font.size = Pt(14)
    r.font.italic = True
    r.font.color.rgb = COLOR_GRAY

    if caption:
        cap = slide.shapes.add_textbox(
            left, top + height + Inches(0.05), width, Inches(0.4)
        )
        ctf = cap.text_frame
        ctf.word_wrap = True
        cp = ctf.paragraphs[0]
        cp.alignment = PP_ALIGN.CENTER
        cr = cp.add_run()
        cr.text = caption
        cr.font.name = FONT
        cr.font.size = Pt(14)
        cr.font.italic = True
        cr.font.color.rgb = COLOR_GRAY


def add_footer(slide, num, total=10):
    """Номер слайда + колонтитул."""
    tb = slide.shapes.add_textbox(
        Inches(0.4), Inches(7.05), Inches(6), Inches(0.35)
    )
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "Нефтепереработка России · XXI век"
    r.font.name = FONT
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_GRAY

    tb2 = slide.shapes.add_textbox(
        Inches(12.0), Inches(7.05), Inches(1.0), Inches(0.35)
    )
    p2 = tb2.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r2 = p2.add_run()
    r2.text = f"{num} / {total}"
    r2.font.name = FONT
    r2.font.size = Pt(11)
    r2.font.color.rgb = COLOR_GRAY


def add_content_slide(prs, title, num):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # пустой layout
    add_title_bar(s, title, prs)
    add_footer(s, num)
    return s


# ----------------------------------------------------------------------
# ПОСТРОЕНИЕ ПРЕЗЕНТАЦИИ
# ----------------------------------------------------------------------
def build():
    prs = Presentation()
    prs.slide_width  = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # ---------- Слайд 1: Титульный ----------
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = COLOR_DARK; bg.line.fill.background()

    # золотая полоса
    line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(3.9),
                              Inches(3.0), Inches(0.08))
    line.fill.solid(); line.fill.fore_color.rgb = COLOR_ACCENT; line.line.fill.background()

    tb = s.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11), Inches(2.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "XXI век: модернизация и новые стандарты"
    r.font.name = FONT; r.font.size = Pt(44); r.font.bold = True
    r.font.color.rgb = COLOR_WHITE

    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "Развитие нефтеперерабатывающей промышленности в России"
    r2.font.name = FONT; r2.font.size = Pt(22); r2.font.color.rgb = COLOR_ACCENT

    # место под фото
    add_image_placeholder(
        s, Inches(8.2), Inches(4.3), Inches(4.3), Inches(2.5),
        caption="Рис. 1. Современный комплекс нефтепереработки."
    )

    tb_meta = s.shapes.add_textbox(Inches(1.0), Inches(4.3), Inches(6.5), Inches(2.0))
    tfm = tb_meta.text_frame; tfm.word_wrap = True
    for txt in ["Выполнил(а): ______________________",
                "Группа: ______________________",
                "2026 г."]:
        p = tfm.paragraphs[0] if txt.startswith("Выполнил") else tfm.add_paragraph()
        r = p.add_run(); r.text = txt
        r.font.name = FONT; r.font.size = Pt(16); r.font.color.rgb = COLOR_WHITE

    # ---------- Слайд 2: Введение: вызовы XXI века ----------
    s = add_content_slide(prs, "Введение: вызовы XXI века", 2)
    add_bullets(s, [
        "Ужесточение экологических норм.",
        "Необходимость повышения глубины переработки нефти.",
        "Переход к выпуску топлива стандартов Евро-5 и выше.",
    ], Inches(0.6), Inches(1.5), Inches(6.0), Inches(4.5), size=20)
    add_image_placeholder(
        s, Inches(7.0), Inches(1.5), Inches(5.8), Inches(4.2),
        caption="Рис. 2. Сравнение стандартов топлива (Евро-3 / Евро-4 / Евро-5)."
    )

    # ---------- Слайд 3: Главная цель модернизации ----------
    s = add_content_slide(prs, "Главная цель модернизации", 3)
    add_bullets(s, [
        "Глубина переработки — ключевой показатель эффективности.",
        "В XXI веке отрасль стремится к безотходному производству.",
    ], Inches(0.6), Inches(1.6), Inches(6.0), Inches(3.0), size=20)
    add_image_placeholder(
        s, Inches(7.0), Inches(1.5), Inches(5.8), Inches(4.2),
        caption="Рис. 3. Динамика роста глубины переработки нефти (70% → 90%)."
    )

    # ---------- Слайд 4: Экологические требования ----------
    s = add_content_slide(prs, "Экологические требования", 4)
    add_bullets(s, [
        "Снижение выбросов CO₂ и SO₂.",
        "Производство экологически чистого бензина и дизеля.",
        "Модернизация очистных сооружений.",
    ], Inches(0.6), Inches(1.6), Inches(6.0), Inches(4.0), size=20)
    add_image_placeholder(
        s, Inches(7.0), Inches(1.5), Inches(5.8), Inches(4.2),
        caption="Рис. 4. Схема «Зелёный НПЗ»: замкнутый цикл водоснабжения."
    )

    # ---------- Слайд 5: Гидроочистка ----------
    s = add_content_slide(prs, "Современные процессы: гидроочистка", 5)
    add_bullets(s, [
        "Удаление серы, азота и кислорода из топлива с помощью водорода.",
        "Ключевой процесс для перехода на стандарт Евро-5.",
    ], Inches(0.6), Inches(1.6), Inches(6.0), Inches(3.0), size=20)
    add_image_placeholder(
        s, Inches(7.0), Inches(1.5), Inches(5.8), Inches(4.2),
        caption="Рис. 5. Принципиальная схема установки гидроочистки."
    )

    # ---------- Слайд 6: Каталитический крекинг ----------
    s = add_content_slide(prs, "Современные процессы: каталитический крекинг", 6)
    add_bullets(s, [
        "Углублённая переработка тяжёлых фракций.",
        "Получение высокооктановых бензинов.",
    ], Inches(0.6), Inches(1.6), Inches(6.0), Inches(3.0), size=20)
    add_image_placeholder(
        s, Inches(7.0), Inches(1.5), Inches(5.8), Inches(4.2),
        caption="Рис. 6. Установка каталитического крекинга (3D-модель)."
    )

    # ---------- Слайд 7: Инновационные технологии ----------
    s = add_content_slide(prs, "Инновационные технологии углубления переработки", 7)
    add_bullets(s, [
        "Гидрокрекинг тяжёлых фракций.",
        "Висбрекинг и коксование остатков.",
        "Переработка мазута и гудрона в светлые нефтепродукты.",
    ], Inches(0.6), Inches(1.6), Inches(6.0), Inches(4.0), size=20)
    add_image_placeholder(
        s, Inches(7.0), Inches(1.5), Inches(5.8), Inches(4.2),
        caption="Рис. 7. Схема переработки тяжёлых нефтяных остатков."
    )

    # ---------- Слайд 8: Цифровизация ----------
    s = add_content_slide(prs, "Цифровизация и автоматизация", 8)
    add_bullets(s, [
        "Внедрение АСУ ТП.",
        "Цифровые двойники технологических установок.",
        "Системы предиктивной аналитики — безопасность и эффективность.",
    ], Inches(0.6), Inches(1.6), Inches(6.0), Inches(4.0), size=20)
    add_image_placeholder(
        s, Inches(7.0), Inches(1.5), Inches(5.8), Inches(4.2),
        caption="Рис. 8. Центр управления современным НПЗ."
    )

    # ---------- Слайд 9: Новые стандарты и продукция ----------
    s = add_content_slide(prs, "Новые стандарты и продукция", 9)
    add_bullets(s, [
        "Топливо стандарта Евро-5 и выше.",
        "Современные масла и битумы.",
        "Интеграция с нефтехимией: производство полимеров.",
    ], Inches(0.6), Inches(1.6), Inches(6.0), Inches(4.0), size=20)
    add_image_placeholder(
        s, Inches(7.0), Inches(1.5), Inches(5.8), Inches(4.2),
        caption="Рис. 9. Продукция: канистра Евро-5, гранулы пластика, дорожное покрытие."
    )

    # ---------- Слайд 10: Заключение ----------
    s = add_content_slide(prs, "Заключение", 10)
    add_bullets(s, [
        "Модернизация НПЗ — залог энергетической безопасности России.",
        "Экологическое благополучие — приоритет XXI века.",
    ], Inches(0.6), Inches(1.6), Inches(6.0), Inches(3.0), size=22)
    add_image_placeholder(
        s, Inches(7.0), Inches(1.5), Inches(5.8), Inches(4.2),
        caption="Рис. 10. Карта России с крупными модернизированными НПЗ."
    )

    prs.save("neftepererabotka_XXI_vek.pptx")
    print("✅ Готово: neftepererabotka_XXI_vek.pptx")


if __name__ == "__main__":
    build()